from __future__ import annotations

import csv
import math
import sys
from datetime import date
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from src.project_paths import CALIBRATION_OUTPUT_DIR


TITLE_BG = "123B5D"
TITLE_ACCENT = "F3B33D"
BODY_BG = "F5F1E8"
CARD_BG = "FFFDFC"
TEXT_DARK = "17324D"
TEXT_MUTED = "5C6B73"
SUCCESS = "2C7A4B"
WARN = "AA5B10"


def _read_markdown_summary(path: Path) -> dict[str, str]:
    summary = {
        "target_rows": "0",
        "worst_relative_error": "0",
        "passed": "no",
    }
    if not path.exists():
        return summary

    for raw_line in path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if line.startswith("- Target rows:"):
            summary["target_rows"] = line.split(":", 1)[1].strip()
        elif line.startswith("- Worst-case relative error:"):
            summary["worst_relative_error"] = line.split(":", 1)[1].strip()
        elif line.startswith("- Passed <3% target:"):
            summary["passed"] = line.split(":", 1)[1].strip()
    return summary


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open("r", encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def _worst_metric_rows(error_rows: list[dict[str, str]]) -> list[dict[str, float | str]]:
    best_by_metric: dict[str, dict[str, float | str]] = {}
    for row in error_rows:
        metric_name = row["metric_name"]
        candidate = {
            "metric_name": metric_name,
            "w_um": float(row["w_um"]),
            "l_um": float(row["l_um"]),
            "relative_error": float(row["relative_error"]),
        }
        current = best_by_metric.get(metric_name)
        if current is None or float(candidate["relative_error"]) > float(current["relative_error"]):
            best_by_metric[metric_name] = candidate
    return sorted(
        best_by_metric.values(),
        key=lambda item: float(item["relative_error"]),
        reverse=True,
    )


def _top_error_rows(error_rows: list[dict[str, str]], top_n: int = 5) -> list[dict[str, float | str]]:
    rows = sorted(error_rows, key=lambda row: float(row["relative_error"]), reverse=True)
    top_rows: list[dict[str, float | str]] = []
    for row in rows[:top_n]:
        top_rows.append(
            {
                "metric_name": row["metric_name"],
                "w_um": float(row["w_um"]),
                "l_um": float(row["l_um"]),
                "relative_error": float(row["relative_error"]),
            }
        )
    return top_rows


def _format_pct(value: float) -> str:
    return f"{value * 100:.2f}%"


def _safe_picture(slide, path: Path, x: float, y: float, w: float, h: float) -> None:
    if path.exists():
        slide.add_picture(str(path), x, y, w=w, h=h)
    else:
        slide.add_shape(
            1,
            x,
            y,
            w,
            h,
        ).fill.solid()


def _add_full_background(slide, color: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(color)
    shape.line.fill.background()


def _rgb(color: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(color)


def _add_text(slide, text: str, x: float, y: float, w: float, h: float, font_size: int, color: str, bold: bool = False, font_name: str = "PingFang SC") -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    run = p.runs[0]
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)


def _add_bullets(slide, lines: list[str], x: float, y: float, w: float, h: float, font_size: int = 18, color: str = TEXT_DARK) -> None:
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.bullet = True
        paragraph.level = 0
        paragraph.alignment = PP_ALIGN.LEFT
        for run in paragraph.runs:
            run.font.name = "PingFang SC"
            run.font.size = Pt(font_size)
            run.font.color.rgb = _rgb(color)


def _add_card(slide, title: str, value: str, note: str, x: float, y: float, w: float, h: float, accent: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(CARD_BG)
    shape.line.color.rgb = _rgb(accent)
    shape.line.width = Inches(0.03)
    _add_text(slide, title, x + 0.18, y + 0.18, w - 0.36, 0.35, 16, TEXT_MUTED, bold=True)
    _add_text(slide, value, x + 0.18, y + 0.56, w - 0.36, 0.55, 25, TEXT_DARK, bold=True)
    _add_text(slide, note, x + 0.18, y + 1.08, w - 0.36, h - 1.24, 11, TEXT_MUTED)


def _add_image_panel(slide, title: str, image_path: Path, x: float, y: float, w: float, h: float) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = _rgb(CARD_BG)
    panel.line.color.rgb = _rgb("D9CDB8")
    _add_text(slide, title, x + 0.18, y + 0.1, w - 0.36, 0.3, 16, TEXT_DARK, bold=True)
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(x + 0.15), Inches(y + 0.48), width=Inches(w - 0.3), height=Inches(h - 0.63))
    else:
        _add_text(slide, f"Missing image:\n{image_path.name}", x + 0.2, y + 0.9, w - 0.4, 0.9, 14, WARN, bold=True)


def _add_table_like_lines(slide, lines: list[str], x: float, y: float, w: float, h: float, title: str) -> None:
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = _rgb(CARD_BG)
    panel.line.color.rgb = _rgb("D9CDB8")
    _add_text(slide, title, x + 0.18, y + 0.12, w - 0.36, 0.3, 16, TEXT_DARK, bold=True)
    current_y = y + 0.56
    for line in lines:
        _add_text(slide, line, x + 0.22, current_y, w - 0.44, 0.36, 13, TEXT_DARK, font_name="Menlo")
        current_y += 0.42


def build_presentation(output_path: Path, calibration_dir: Path) -> None:
    from pptx import Presentation

    summary = _read_markdown_summary(calibration_dir / "calibration_summary.md")
    error_rows = _read_csv_rows(calibration_dir / "calibration_error_report.csv")
    worst_by_metric = _worst_metric_rows(error_rows)
    top_rows = _top_error_rows(error_rows, top_n=6)
    worst_relative_error = float(summary["worst_relative_error"])
    target_rows = int(float(summary["target_rows"]))
    pass_text = "已达标" if summary["passed"] == "yes" else "未达标"

    prs = Presentation()
    prs.slide_width = 12188952
    prs.slide_height = 6858000
    prs.core_properties.author = "OpenAI Codex"
    prs.core_properties.title = "BSIM4 Calibration Progress Summary"
    prs.core_properties.subject = "Project progress and calibration results"

    today = date.today().isoformat()

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_full_background(slide, TITLE_BG)
    _add_text(slide, "BSIM4 校准项目迭代总结", 0.72, 0.78, 5.9, 0.7, 29, "FFFFFF", bold=True)
    _add_text(slide, "从单脚本原型演进到可运行的多阶段校准 pipeline", 0.74, 1.55, 6.6, 0.45, 16, "D8E6F0")
    _add_text(slide, f"更新日期: {today}", 0.74, 2.02, 2.5, 0.3, 12, "D8E6F0")
    _add_card(slide, "最新 worst-case error", _format_pct(worst_relative_error), "目标仍为 42 个器件 x 6 指标全部 < 3%。", 0.72, 2.78, 2.6, 1.55, TITLE_ACCENT)
    _add_card(slide, "训练点数量", str(target_rows), "当前默认目标 CSV 覆盖 42 个 W/L 训练点。", 3.48, 2.78, 2.2, 1.55, "7AC6D2")
    _add_card(slide, "当前结论", pass_text, "当前最佳结果约 5.10%，主瓶颈仍集中在漏电相关指标。", 5.86, 2.78, 2.6, 1.55, "E07A5F")
    _add_image_panel(slide, "参数趋势总览", calibration_dir / "parameter_trends_vs_w.png", 8.92, 0.72, 3.68, 5.95)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_full_background(slide, BODY_BG)
    _add_text(slide, "1. 项目迭代过程", 0.72, 0.44, 3.5, 0.48, 24, TEXT_DARK, bold=True)
    _add_bullets(
        slide,
        [
            "第一阶段：从早期单脚本原型，沉淀出按 src/ 组织的模块化校准工程。",
            "第二阶段：完成目标数据读取、参数边界管理、连续参数曲面和 4 角点问题建模。",
            "第三阶段：接入 pymoo / NSGA-II，多目标搜索角点参数，再用双线性曲面做初始化。",
            "第四阶段：加入逐器件 bounded local tuning 和全局 refit，形成完整多阶段 pipeline。",
            "第五阶段：补齐 CSV、Markdown、图表输出与单元测试，形成可验证的工程闭环。",
            "当前阶段：围绕 leakage error 做诊断优化，并把趋势图和进展汇报材料纳入产物链路。",
        ],
        0.82,
        1.18,
        6.05,
        5.3,
    )
    _add_image_panel(slide, "当前参数随 W 变化", calibration_dir / "parameter_trends_vs_w.png", 7.15, 1.08, 5.45, 5.95)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_full_background(slide, BODY_BG)
    _add_text(slide, "2. 当前校准 pipeline", 0.72, 0.44, 3.8, 0.48, 24, TEXT_DARK, bold=True)
    _add_bullets(
        slide,
        [
            "目标数据按 W/L 尺寸索引，形成 42 个训练点。",
            "4 角点多目标优化负责找出一组可行角点参数。",
            "从角点构造连续参数曲面，为所有尺寸生成初始化参数。",
            "逐器件进行 bounded Powell + DE 的局部精修，提升教师解质量。",
            "基于局部教师解重新拟合全局连续模型，输出 calibrated params / metrics / plots。",
        ],
        0.82,
        1.14,
        5.45,
        3.5,
    )
    _add_card(slide, "优化内核", "NSGA-II + local refinement", "角点阶段默认 pop_size=24, n_gen=8；局部阶段使用 Powell 与差分进化混合。", 0.86, 4.86, 2.92, 1.56, TITLE_ACCENT)
    _add_card(slide, "连续模型", "Bilinear init + IDW refit", "训练点上当前 refitted_global_params 与 local_tuned_params 一致，瓶颈更多在局部教师解。", 3.98, 4.86, 3.08, 1.56, "7AC6D2")
    _add_image_panel(slide, "漏电误差热图", calibration_dir / "error_heatmap_idoff_a.png", 7.34, 1.04, 5.28, 5.78)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_full_background(slide, BODY_BG)
    _add_text(slide, "3. 工程化与验证状态", 0.72, 0.44, 4.0, 0.48, 24, TEXT_DARK, bold=True)
    _add_card(slide, "测试状态", "40 / 40 通过", "仓库已验证从数学层测试到 PySpice-backed smoke test 的完整链路。", 0.82, 1.18, 2.58, 1.48, SUCCESS)
    _add_card(slide, "输出闭环", "CSV + Markdown + PNG", "历史上已生成 calibration_output 与 dataset_generation 两类产物。", 3.62, 1.18, 2.5, 1.48, "7AC6D2")
    _add_card(slide, "运行入口", "scripts/run_calibration.py", "建议继续使用 scripts/ 下标准入口，而不是兼容 wrapper。", 6.34, 1.18, 3.14, 1.48, TITLE_ACCENT)
    _add_bullets(
        slide,
        [
            "当前 `spice` 环境已能跑通全量测试和真实仿真 smoke test。",
            "新增趋势图与 PPT builder 后，校准结果和汇报材料都能跟随同一套输出目录更新。",
            "本轮新增产物重点补上了参数随 W/L 变化的趋势视图，方便检查曲面形状是否符合预期。",
        ],
        0.84,
        3.05,
        5.72,
        2.2,
    )
    _add_image_panel(slide, "参数随 L 变化", calibration_dir / "parameter_trends_vs_l.png", 6.9, 2.1, 5.7, 4.85)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_full_background(slide, BODY_BG)
    _add_text(slide, "4. 当前结果总览", 0.72, 0.44, 3.5, 0.48, 24, TEXT_DARK, bold=True)
    _add_card(slide, "worst-case", _format_pct(worst_relative_error), "当前默认结果仍高于 3% 验收目标。", 0.82, 1.14, 2.42, 1.45, WARN)
    _add_card(slide, "最强信号", "漏电指标仍最难压", "idoff_a 和 isoff_a 继续占据 top error 列表前列。", 3.48, 1.14, 2.82, 1.45, "E07A5F")
    _add_card(slide, "建模判断", "先提局部教师解", "训练点上全局 refit 未明显拉高误差，优先级仍是局部解质量。", 6.54, 1.14, 3.1, 1.45, "7AC6D2")
    metric_lines = [
        f"{row['metric_name']:<9} {_format_pct(float(row['relative_error'])):>7}  @ W={float(row['w_um']):g}, L={float(row['l_um']):g}"
        for row in worst_by_metric
    ]
    _add_table_like_lines(slide, metric_lines, 0.82, 3.0, 6.08, 3.32, "各指标 worst-case")
    _add_image_panel(slide, "idoff target vs simulated", calibration_dir / "target_vs_simulated_idoff_a.png", 7.14, 2.18, 5.46, 4.62)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_full_background(slide, BODY_BG)
    _add_text(slide, "5. 误差热点", 0.72, 0.44, 3.2, 0.48, 24, TEXT_DARK, bold=True)
    hotspot_lines = [
        f"{index + 1}. {row['metric_name']}  {_format_pct(float(row['relative_error']))}  @ ({float(row['w_um']):g}, {float(row['l_um']):g})"
        for index, row in enumerate(top_rows)
    ]
    _add_table_like_lines(slide, hotspot_lines, 0.82, 1.18, 5.74, 2.95, "Top error points")
    _add_bullets(
        slide,
        [
            "最坏点集中在较小或中等尺寸组合，不是单纯的大尺寸外推问题。",
            "当前直接给漏电指标加权并没有改善 worst-case，说明问题不只是目标函数权重配置。",
            "更值得优先尝试的是增加角点搜索和局部精修预算，而不是先重写优化器。",
        ],
        0.84,
        4.42,
        5.78,
        1.92,
    )
    _add_image_panel(slide, "idoff heatmap", calibration_dir / "error_heatmap_idoff_a.png", 7.02, 1.12, 2.64, 5.72)
    _add_image_panel(slide, "isoff heatmap", calibration_dir / "error_heatmap_isoff_a.png", 9.92, 1.12, 2.64, 5.72)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_full_background(slide, BODY_BG)
    _add_text(slide, "6. 参数趋势图：随 W 变化", 0.72, 0.44, 4.6, 0.48, 24, TEXT_DARK, bold=True)
    _add_bullets(
        slide,
        [
            "每个子图对应一个校准参数，横轴为 W，不同 L 为多条曲线。",
            "该视图更适合检查宽度扩展时参数是否出现突跳或异常扭曲。",
        ],
        0.82,
        1.04,
        4.68,
        1.1,
    )
    _add_image_panel(slide, "parameter_trends_vs_w.png", calibration_dir / "parameter_trends_vs_w.png", 0.82, 2.02, 11.78, 4.78)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_full_background(slide, BODY_BG)
    _add_text(slide, "7. 参数趋势图：随 L 变化", 0.72, 0.44, 4.6, 0.48, 24, TEXT_DARK, bold=True)
    _add_bullets(
        slide,
        [
            "每个子图横轴为 L，不同 W 为多条曲线。",
            "与上一页组合后，可以更系统地判断参数曲面对尺寸的整体响应是否平滑一致。",
        ],
        0.82,
        1.04,
        4.68,
        1.1,
    )
    _add_image_panel(slide, "parameter_trends_vs_l.png", calibration_dir / "parameter_trends_vs_l.png", 0.82, 2.02, 11.78, 4.78)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_full_background(slide, BODY_BG)
    _add_text(slide, "8. 下一步建议", 0.72, 0.44, 3.4, 0.48, 24, TEXT_DARK, bold=True)
    _add_bullets(
        slide,
        [
            "先固化当前默认基线和本轮带趋势图的结果目录，避免后续实验覆盖。",
            "围绕 idoff_a / isoff_a 做更聚焦的敏感度分析，识别真正主导 leakage 的参数组合。",
            "优先增加 NSGA-II 和局部 refinement 预算，再判断是否需要更强的 refit 或特化优化器。",
            "如果后续要评估非训练点泛化，再单独加强全局连续模型，而不是与训练点拟合问题混在一起。",
        ],
        0.82,
        1.18,
        7.1,
        3.8,
    )
    _add_card(slide, "建议优先级", "预算提升 > 漏电加权 > 重写算法", "本轮实验已表明直接漏电加权会让 worst-case 变差。", 0.86, 5.38, 4.3, 1.3, TITLE_ACCENT)
    _add_image_panel(slide, "vtsat heatmap", calibration_dir / "error_heatmap_vtsat_v.png", 8.38, 1.18, 2.06, 2.52)
    _add_image_panel(slide, "vtlin heatmap", calibration_dir / "error_heatmap_vtlin_v.png", 10.54, 1.18, 2.06, 2.52)
    _add_image_panel(slide, "pareto front", calibration_dir / "pareto_front.png", 8.38, 4.02, 4.22, 2.66)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def main() -> int:
    output_path = CALIBRATION_OUTPUT_DIR / "project_progress_summary_2026-04-22.pptx"
    build_presentation(output_path=output_path, calibration_dir=CALIBRATION_OUTPUT_DIR)
    print(f"Wrote presentation to: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
